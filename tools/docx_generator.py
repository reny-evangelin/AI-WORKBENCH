from pathlib import Path

from docx import Document
from docx.shared import Pt


def generate_docx(data: dict, output_path: str) -> str:
    """
    Generate an Inspection Approval Note in DOCX format.

    Args:
        data: Structured analysis data.
        output_path: Path where the DOCX file should be created.

    Returns:
        The output file path.
    """

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    # Title
    title = document.add_heading(
        data.get("title", "Engineering Report"),
        level=0
    )
    title.alignment = 1

    # Backward compatibility for old equipment table
    if "equipment" in data:
        document.add_heading("Equipment Details", level=1)
        equipment_table = document.add_table(rows=0, cols=2)
        equipment_table.style = "Table Grid"

        equipment_details = [
            ("Equipment", data.get("equipment", "")),
            ("Equipment Type", data.get("equipment_type", "")),
            ("Location", data.get("location", "")),
            ("Severity", data.get("severity", "")),
        ]

        for label, value in equipment_details:
            row = equipment_table.add_row().cells
            row[0].text = label
            row[1].text = str(value)

    # Dynamic sections
    sections = data.get("sections", [])
    if sections:
        for sec in sections:
            document.add_heading(str(sec.get("heading", "")), level=1)
            document.add_paragraph(str(sec.get("content", "")))
    else:
        # Fallback for old schema
        if "finding" in data:
            document.add_heading("Inspection Finding", level=1)
            document.add_paragraph(data.get("finding", ""))
        if "analysis" in data:
            document.add_heading("AI Analysis", level=1)
            document.add_paragraph(data.get("analysis", ""))
        if "recommendation" in data:
            document.add_heading("Recommended Action", level=1)
            document.add_paragraph(data.get("recommendation", ""))
        if "sop_reference" in data:
            document.add_heading("SOP / Knowledge Reference", level=1)
            document.add_paragraph(data.get("sop_reference", ""))

    # Source Traceability
    if "source_document" in data:
        document.add_heading("Source Traceability", level=1)
        source_document = data.get("source_document", "")
        source_pages = data.get("source_pages", [])
        document.add_paragraph(f"Source Document: {source_document}")
        document.add_paragraph(f"Source Pages: {', '.join(map(str, source_pages))}")

    # Basic font formatting
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10)

    document.save(output_file)

    return str(output_file)