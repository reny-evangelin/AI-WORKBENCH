import sys
import os
import json
from pathlib import Path
from agent import run_agent

def verify_file(filepath: str, ext: str):
    p = Path(filepath)
    if not p.exists():
        print(f"[ERROR] File does not exist: {filepath}")
        return False
    if p.stat().st_size == 0:
        print(f"[ERROR] File is empty: {filepath}")
        return False
    if p.suffix != ext:
        print(f"[ERROR] Wrong extension {p.suffix} expected {ext}")
        return False
    
    if ext == ".xlsx":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(p)
            print(f"[OK] Workbook opened successfully. Sheets: {wb.sheetnames}")
            # Read first sheet
            ws = wb.active
            text = " ".join([str(c) for r in ws.iter_rows(values_only=True) for c in r if c])
            print(f"[DEBUG] Excel first sheet content dump length: {len(text)}")
        except Exception as e:
            print(f"[ERROR] Failed to open workbook: {e}")
            return False
            
    elif ext == ".pdf":
        import fitz
        try:
            doc = fitz.open(p)
            print(f"[OK] PDF opened successfully. Pages: {len(doc)}")
            text = " ".join([p.get_text() for p in doc])
            print(f"[DEBUG] PDF content dump length: {len(text)}")
            if "comment" not in text.lower():
                print("[WARNING] 'comment' not found in PDF!")
            if "recommendation" not in text.lower():
                print("[WARNING] 'recommendation' not found in PDF!")
        except Exception as e:
            print(f"[ERROR] Failed to open PDF: {e}")
            return False
            
    return True

print("========== TEST 1: EXCEL MULTI-SHEET ==========")
req1 = (
    "Create an Excel report for our AI project team.\n\n"
    "Create sheets:\n"
    "Team\n"
    "Tasks\n"
    "Comments\n\n"
    "The Team sheet should contain:\n"
    "Name\n"
    "Role\n"
    "Module\n"
    "Status\n\n"
    "The Tasks sheet should contain:\n"
    "Task\n"
    "Member\n"
    "Status\n\n"
    "The Comments sheet should contain:\n"
    "Member\n"
    "Comment"
)

res1 = run_agent(req1)
print(f"Status: {res1.status}")
print(f"Response: {res1.answer}")
try:
    d1 = json.loads(res1.answer)
    print(f"File Path: {d1.get('path')}")
    verify_file(d1.get("path"), ".xlsx")
except Exception as e:
    print(f"[ERROR] Failed to parse response or verify file: {e}")

print("\n========== TEST 2: PDF WITH COMMENTS/RECOMMENDATIONS ==========")
req2 = "Create a PDF report of the same project with comments and recommendations."
res2 = run_agent(req2)
print(f"Status: {res2.status}")
print(f"Response: {res2.answer}")
try:
    d2 = json.loads(res2.answer)
    print(f"File Path: {d2.get('path')}")
    verify_file(d2.get("path"), ".pdf")
except Exception as e:
    print(f"[ERROR] Failed to parse response or verify file: {e}")
