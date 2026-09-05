import pytest
import os
import json
from pathlib import Path
from agent.graph.nodes import understand_request
from tools.excel_generator import generate_excel
from tools.output_validator import validate_analysis_data
from agent.graph.state import AgentState
from agent.schemas import ExcelPlan
import openpyxl

@pytest.fixture
def test_output_dir(tmp_path):
    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    return str(output_dir)

def test_1_basic_excel_schema_validation():
    """Test 1 - Valid schema is accepted."""
    data = {
        "intent": "excel_generation",
        "filename": "student_report.xlsx",
        "title": "Students",
        "sheets": [
            {
                "name": "Students",
                "columns": ["Name", "Marks"],
                "rows": [{"Name": "Arun", "Marks": 90}],
                "formulas": []
            }
        ]
    }
    is_valid, errors = validate_analysis_data(data)
    assert is_valid, f"Expected valid data, got errors: {errors}"

def test_2_empty_template_excel(test_output_dir):
    """Test 2 - Empty template generation."""
    data = {
        "intent": "excel_generation",
        "filename": "template.xlsx",
        "title": "Template",
        "sheets": [
            {
                "name": "TemplateSheet",
                "columns": ["ID", "Name", "Score"],
                "rows": [],
                "formulas": []
            }
        ]
    }
    path = generate_excel(data, test_output_dir)
    assert os.path.exists(path)
    
    wb = openpyxl.load_workbook(path)
    assert "TemplateSheet" in wb.sheetnames
    ws = wb["TemplateSheet"]
    
    # Headers should be written
    assert ws["A2"].value == "ID"
    assert ws["B2"].value == "Name"
    assert ws["C2"].value == "Score"
    
    # Row 3 should be empty
    assert ws["A3"].value is None

def test_3_multiple_sheets(test_output_dir):
    """Test 3 - Multiple sheets generation."""
    data = {
        "intent": "excel_generation",
        "filename": "multi.xlsx",
        "title": "Multi",
        "sheets": [
            {
                "name": "SheetOne",
                "columns": ["A"],
                "rows": [{"A": "1"}]
            },
            {
                "name": "SheetTwo",
                "columns": ["B"],
                "rows": [{"B": "2"}]
            }
        ]
    }
    path = generate_excel(data, test_output_dir)
    wb = openpyxl.load_workbook(path)
    assert "SheetOne" in wb.sheetnames
    assert "SheetTwo" in wb.sheetnames

def test_4_formulas(test_output_dir):
    """Test 4 - Formula generation and injection."""
    data = {
        "intent": "excel_generation",
        "filename": "calc.xlsx",
        "title": "Calc",
        "sheets": [
            {
                "name": "Maths",
                "columns": ["Val1", "Val2", "Total"],
                "rows": [
                    {"Val1": 10, "Val2": 20}
                ],
                "formulas": [
                    {"column": "Total", "formula": "=SUM(A{row}:B{row})"}
                ]
            }
        ]
    }
    path = generate_excel(data, test_output_dir)
    wb = openpyxl.load_workbook(path)
    ws = wb["Maths"]
    assert ws["C3"].value == "=SUM(A3:B3)"

def test_7_invalid_schema():
    """Test 7 - Invalid schema rejected."""
    data = {
        "intent": "excel_generation",
        "filename": "bad.xlsx",
        "sheets": [
            {
                "name": "BadSheet",
                "columns": ["A", "A"], # duplicate column
                "rows": [{"A": 1}]
            }
        ]
    }
    is_valid, errors = validate_analysis_data(data)
    assert not is_valid
    assert any("duplicate columns" in e for e in errors)
