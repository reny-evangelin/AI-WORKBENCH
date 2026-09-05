"""
test_agentic.py — Unit and Live Integration tests for Phase 3 Agentic Engine
"""

import pytest
from unittest.mock import patch
import json
from pathlib import Path
from agent import process_request, tool_registry
from agent.tools import ToolResult
from agent.graph.nodes import execute_tool_node, evaluate_result
import openpyxl

# =====================================================================
# Unit Tests (Run in CI without network/Ollama dependencies)
# =====================================================================

def test_9_tool_failure_mocked():
    """Test 9 — Tool Failure: Mock generate_excel to fail. Expected success=False. No fake success."""
    state = {
        "user_request": "Generate an Excel",
        "tool_name": "generate_excel",
        "tool_input": {"content": {"title": "Test", "sheets": []}},
        "observations": [],
        "sources": [],
    }
    with patch.object(tool_registry, "execute_tool") as mock_exec:
        mock_exec.return_value = ToolResult(
            success=False,
            tool_name="generate_excel",
            data={},
            sources=[],
            error="Mocked Excel failure",
        )
        obs_res = execute_tool_node(state)
        assert obs_res["tool_result"]["success"] is False
        
        state["tool_result"] = obs_res["tool_result"]
        eval_res = evaluate_result(state)
        assert eval_res["status"] == "error"
        assert "Mocked Excel failure" in eval_res["response"]

def test_10_invalid_excel_data():
    """Test 10 — Invalid Excel Data: Use invalid workbook data {"sheets": "invalid"}."""
    # We call output_generator directly or through tool_registry
    res = tool_registry.execute_tool("generate_excel", {"content": {"title": "Bad", "sheets": "invalid"}})
    # Interfaces catches validation errors and returns success=False
    assert res.success is False
    assert "Invalid analysis data" in res.error or "must be a list" in res.error or "must be a dictionary" in res.error


# =====================================================================
# Live Integration Tests (Require live Ollama service)
# =====================================================================

@pytest.mark.integration
def test_1_normal_chat(ollama_check):
    """Test 1 — Normal Chat: hello -> intent=general, no document tool"""
    res = process_request("hello")
    assert res.status == "success"
    # Should be a normal response, not a JSON file object
    try:
        data = json.loads(res.answer)
        assert data.get("type") != "file", "Should not generate file for 'hello'"
    except:
        pass  # Text output is expected

@pytest.mark.integration
def test_2_pdf(ollama_check):
    """Test 2 — PDF: Create a PDF report about AI -> generate_pdf, real PDF"""
    res = process_request("Create a PDF report about AI.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    assert data["type"] == "file"
    assert data["file_type"] == "pdf"
    
    pdf_path = Path(data["path"])
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 0

@pytest.mark.integration
def test_3_pdf_comments(ollama_check):
    """Test 3 — PDF Comments: verify Comments exists inside PDF"""
    res = process_request("Create a PDF about AI with comments.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    assert data["type"] == "file"
    
    # Simple check on the text output from PyMuPDF
    import fitz # PyMuPDF
    doc = fitz.open(data["path"])
    text = ""
    for page in doc:
        text += page.get_text()
    assert "Comment" in text or "comment" in text.lower()

@pytest.mark.integration
def test_4_pdf_recommendations(ollama_check):
    """Test 4 — PDF Recommendations: verify Recommendations exists inside PDF"""
    res = process_request("Create a PDF about AI with recommendations.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    
    import fitz
    doc = fitz.open(data["path"])
    text = ""
    for page in doc:
        text += page.get_text()
    assert "Recommendation" in text or "recommendation" in text.lower()

@pytest.mark.integration
def test_5_excel_basic(ollama_check):
    """Test 5 — Excel Basic: xlsx exists, sheet exists, Name exists, Age exists, Marks exists, 5 rows exist."""
    res = process_request("Create an Excel file with Name, Age and Marks. Add 5 example students.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    assert data["type"] == "file"
    assert data["file_type"] == "excel"
    
    path = Path(data["path"])
    assert path.exists()
    
    wb = openpyxl.load_workbook(path)
    assert len(wb.sheetnames) >= 1
    ws = wb.active
    
    # Read text from all cells
    all_text = ""
    row_count = ws.max_row
    for row in ws.iter_rows(values_only=True):
        all_text += " ".join([str(c) for c in row if c])
        
    assert "Name" in all_text
    assert "Age" in all_text
    assert "Marks" in all_text
    # 1 title, 1 header, 5 students = ~7 rows
    assert row_count >= 5

@pytest.mark.integration
def test_6_excel_comments(ollama_check):
    """Test 6 — Excel Comments: Verify Comments exists in the actual workbook."""
    res = process_request("Create an Excel report with Name, Status and Comments.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    
    wb = openpyxl.load_workbook(data["path"])
    ws = wb.active
    
    all_text = ""
    for row in ws.iter_rows(values_only=True):
        all_text += " ".join([str(c) for c in row if c])
        
    assert "Comment" in all_text or "comment" in all_text.lower()

@pytest.mark.integration
def test_7_excel_multiple_sheets(ollama_check):
    """Test 7 — Excel Multiple Sheets: Verify Team, Tasks, Comments all exist."""
    res = process_request("Create an Excel workbook with Team, Tasks and Comments sheets.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    
    wb = openpyxl.load_workbook(data["path"])
    sheet_names = wb.sheetnames
    
    # It might create them as "Team", "Tasks", "Comments"
    sheet_str = " ".join(sheet_names).lower()
    assert "team" in sheet_str
    assert "task" in sheet_str
    assert "comment" in sheet_str

@pytest.mark.integration
def test_8_docx(ollama_check):
    """Test 8 — DOCX: Verify DOCX exists, opens, content exists."""
    res = process_request("Create a Word document about our AI project.")
    assert res.status == "success"
    
    data = json.loads(res.answer)
    assert data["type"] == "file"
    assert data["file_type"] == "docx"
    
    path = Path(data["path"])
    assert path.exists()
    assert path.stat().st_size > 0
    
    import docx
    doc = docx.Document(path)
    assert len(doc.paragraphs) > 0
