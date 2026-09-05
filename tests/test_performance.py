"""
tests/test_performance.py
=========================
Performance benchmarks for the Smart Router and input processing pipeline.

Targets:
  SmartRouter decision:  < 5 ms  (deterministic Python, never LLM)
  PDF text extraction:   < 2000 ms
  DOCX extraction:       < 500 ms
  XLSX extraction:       < 500 ms
  TXT read:              < 100 ms

Run with:
  pytest tests/test_performance.py -v
"""

import os
import sys
import time
import tempfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from api.main import SmartRouter, _process_docx, _process_xlsx, _process_txt


# ---------------------------------------------------------------------------
# Fixtures / helpers
# ---------------------------------------------------------------------------

@pytest.fixture
def router():
    return SmartRouter()


def _make_digital_pdf(n_pages: int = 3) -> str:
    """Create a digital PDF with selectable text."""
    try:
        import pymupdf
        doc = pymupdf.open()
        for _ in range(n_pages):
            page = doc.new_page()
            page.insert_text(
                (50, 72),
                "Engineering document content with lots of readable text. " * 10
            )
        fd, path = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        doc.save(path)
        doc.close()
        return path
    except ImportError:
        pytest.skip("pymupdf not installed")


def _make_docx(n_paragraphs: int = 50) -> str:
    try:
        from docx import Document
        doc = Document()
        for i in range(n_paragraphs):
            doc.add_paragraph(f"Paragraph {i}: Equipment TK-{100+i} pressure vessel safety valve.")
        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc.save(path)
        return path
    except ImportError:
        pytest.skip("python-docx not installed")


def _make_xlsx(n_rows: int = 200) -> str:
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Tag", "Type", "Pressure", "Temperature", "Status"])
        for i in range(n_rows):
            ws.append([f"TK-{i:04d}", "Vessel", f"{i*1.2:.1f} bar", f"{20+i*0.1:.1f}°C", "Active"])
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        return path
    except ImportError:
        pytest.skip("openpyxl not installed")


# ---------------------------------------------------------------------------
# Task 8: SmartRouter must be extremely fast (never LLM)
# ---------------------------------------------------------------------------

class TestRouterPerformance:
    """The SmartRouter must always be sub-5ms. Never calls Ollama."""

    ROUTER_MAX_MS = 5.0   # strict: deterministic Python

    def _measure(self, router, message, file_path, filename, input_type, file_type) -> float:
        # Average over 5 runs to avoid cold-start noise
        times = []
        for _ in range(5):
            t0 = time.perf_counter()
            router.route(message, file_path, filename, input_type, file_type)
            times.append((time.perf_counter() - t0) * 1000)
        return sum(times) / len(times)

    def test_text_routing_is_fast(self, router):
        avg_ms = self._measure(router, "Hello", None, None, "text", None)
        print(f"\n  [PERF] text routing: {avg_ms:.3f}ms")
        assert avg_ms < self.ROUTER_MAX_MS, f"Router too slow: {avg_ms:.2f}ms > {self.ROUTER_MAX_MS}ms"

    def test_image_routing_is_fast(self, router):
        fd, path = tempfile.mkstemp(suffix=".png")
        os.write(fd, b"\x89PNG\r\n")
        os.close(fd)
        try:
            avg_ms = self._measure(router, "Explain this", path, "img.png", "image", "png")
            print(f"\n  [PERF] image routing: {avg_ms:.3f}ms")
            assert avg_ms < self.ROUTER_MAX_MS
        finally:
            os.remove(path)

    def test_csv_rejection_is_fast(self, router):
        fd, path = tempfile.mkstemp(suffix=".csv")
        os.write(fd, b"col1,col2\n1,2\n")
        os.close(fd)
        try:
            avg_ms = self._measure(router, "Analyze", path, "data.csv", "csv", "csv")
            print(f"\n  [PERF] csv rejection: {avg_ms:.3f}ms")
            assert avg_ms < self.ROUTER_MAX_MS
        finally:
            os.remove(path)

    def test_pdf_routing_is_fast(self, router):
        """PDF routing includes text density check — still must be fast."""
        path = _make_digital_pdf()
        try:
            avg_ms = self._measure(router, "Summarize", path, "doc.pdf", "pdf", "pdf")
            print(f"\n  [PERF] pdf routing (with density check): {avg_ms:.1f}ms")
            # PDF check reads the file so we allow up to 500ms
            assert avg_ms < 500.0, f"PDF routing too slow: {avg_ms:.1f}ms"
        finally:
            os.remove(path)


# ---------------------------------------------------------------------------
# Task 3: Verify real processing is within acceptable bounds
# ---------------------------------------------------------------------------

class TestProcessingPerformance:
    """Processing helpers must be fast enough for production use."""

    MAX_PDF_MS = 2000
    MAX_DOCX_MS = 500
    MAX_XLSX_MS = 500
    MAX_TXT_MS = 100

    def test_pdf_text_extraction_timing(self):
        path = _make_digital_pdf(n_pages=5)
        try:
            t0 = time.perf_counter()
            from api.main import _process_pdf_digital
            text = _process_pdf_digital(path)
            elapsed = (time.perf_counter() - t0) * 1000
            print(f"\n  [PERF] PDF extraction (5 pages): {elapsed:.1f}ms, {len(text)} chars")
            assert len(text) > 100, "Should extract meaningful text"
            assert elapsed < self.MAX_PDF_MS, f"Too slow: {elapsed:.1f}ms"
        finally:
            os.remove(path)

    def test_docx_extraction_timing(self):
        path = _make_docx(n_paragraphs=100)
        try:
            t0 = time.perf_counter()
            text = _process_docx(path)
            elapsed = (time.perf_counter() - t0) * 1000
            print(f"\n  [PERF] DOCX extraction (100 paras): {elapsed:.1f}ms, {len(text)} chars")
            assert "Paragraph 0" in text
            assert elapsed < self.MAX_DOCX_MS
        finally:
            os.remove(path)

    def test_xlsx_extraction_timing(self):
        path = _make_xlsx(n_rows=200)
        try:
            t0 = time.perf_counter()
            text = _process_xlsx(path)
            elapsed = (time.perf_counter() - t0) * 1000
            print(f"\n  [PERF] XLSX extraction (200 rows): {elapsed:.1f}ms, {len(text)} chars")
            assert "TK-0000" in text
            assert elapsed < self.MAX_XLSX_MS
        finally:
            os.remove(path)

    def test_txt_read_timing(self):
        fd, path = tempfile.mkstemp(suffix=".txt")
        content = ("Engineering safety report. " * 100).encode()
        os.write(fd, content)
        os.close(fd)
        try:
            t0 = time.perf_counter()
            text = _process_txt(path)
            elapsed = (time.perf_counter() - t0) * 1000
            print(f"\n  [PERF] TXT read: {elapsed:.2f}ms, {len(text)} chars")
            assert len(text) > 100
            assert elapsed < self.MAX_TXT_MS
        finally:
            os.remove(path)


# ---------------------------------------------------------------------------
# Task 6: Verify single-parse behaviour (no duplicate processing)
# ---------------------------------------------------------------------------

class TestNoDuplicateProcessing:
    """Verify that file content is extracted exactly once per request."""

    def test_pdf_extracted_once(self):
        """_process_pdf_digital should open the file once and return."""
        path = _make_digital_pdf()
        try:
            from api.main import _process_pdf_digital
            call_count = {"n": 0}
            original_open = __builtins__.__dict__.get("open") if hasattr(__builtins__, "__dict__") else None

            import pymupdf
            original_pymupdf_open = pymupdf.open

            opened = []

            def tracking_open(p, *a, **kw):
                opened.append(p)
                return original_pymupdf_open(p, *a, **kw)

            pymupdf.open = tracking_open
            try:
                _process_pdf_digital(path)
            finally:
                pymupdf.open = original_pymupdf_open

            # Should open exactly once
            assert len(opened) == 1, f"PDF opened {len(opened)} times, expected 1"
        finally:
            os.remove(path)


# ---------------------------------------------------------------------------
# Summary printer
# ---------------------------------------------------------------------------

def test_print_summary():
    """Print a formatted summary of routing performance."""
    router = SmartRouter()
    results = {}

    cases = [
        ("text",     None,   None,      "text",  None),
        ("csv_rej",  None,   "x.csv",   "csv",   "csv"),
    ]

    for label, path, fname, itype, ftype in cases:
        t0 = time.perf_counter()
        for _ in range(10):
            router.route("Test message", path, fname, itype, ftype)
        avg = (time.perf_counter() - t0) * 1000 / 10
        results[label] = avg

    print("\n" + "=" * 50)
    print("  SMART ROUTER PERFORMANCE SUMMARY")
    print("=" * 50)
    for label, ms in results.items():
        status = "PASS" if ms < 5.0 else "WARN"
        print(f"  [{status}] {label:20s}: {ms:.3f}ms")
    print("=" * 50)
    print("  TARGET: < 5ms (deterministic Python, no LLM)")
    print("=" * 50 + "\n")
