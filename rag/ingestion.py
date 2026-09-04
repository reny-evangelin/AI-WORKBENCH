# pyrefly: ignore [missing-import]
import pymupdf
from docx import Document
import openpyxl


def extract_text_from_pdf(file_path):
    document = pymupdf.open(file_path)
    pages = []

    for page_number, page in enumerate(document):
        text = page.get_text()

        if text.strip():
            pages.append({
                "text": text.strip(),
                "page": page_number + 1,
                "source": file_path
            })

    document.close()
    return pages


def extract_text_from_docx(file_path):
    document = Document(file_path)
    sections = []

    for i, para in enumerate(document.paragraphs):
        if para.text.strip():
            sections.append({
                "text": para.text.strip(),
                "section": i + 1,
                "source": file_path
            })

    return sections


def extract_text_from_excel(file_path):
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    rows_out = []

    for sheet in workbook.sheetnames:
        ws = workbook[sheet]
        headers = [cell.value for cell in ws[1]]

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row_text = ", ".join(
                f"{h}: {v}" for h, v in zip(headers, row) if v is not None
            )
            if row_text.strip():
                rows_out.append({
                    "text": row_text,
                    "sheet": sheet,
                    "row": row_number,
                    "source": file_path
                })

    return rows_out